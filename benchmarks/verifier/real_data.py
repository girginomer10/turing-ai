#!/usr/bin/env python3
"""Build a PRIVATE real-message audit set from the owner's own ChatGPT and Claude exports. Runs entirely locally; it never
prints message text. Real messages are written only to a private directory OUTSIDE the repository and must never be committed.

Only the owner's own typed messages are used (ChatGPT role 'user', Claude sender 'human'); assistant text is never read. ChatGPT
conversations flagged 'do not remember' are skipped, as are custom-instruction/context messages, pasted code and long documents.

Three steps (the review sheet is an .xlsx so Excel shows Turkish characters and columns correctly):
  python3 -I real_data.py build     reads the exports, selects whole short runs of messages from a seeded sample of conversations,
                                    masks emails / phones / URLs / IDs / keys / addresses and replaces common first names with
                                    stable pseudonyms, and writes review.csv to the private directory. Capitalised words that
                                    may still be names or places are listed in 'check' for the owner to look at.
  python3 -I real_data.py xlsx      converts review.csv to review.xlsx (UTF-8, wrapped text, keep dropdown 1/0). Needs openpyxl.
  python3 -I real_data.py finalize  reads review.xlsx (or review.csv if no xlsx exists) after the owner has edited it (edit 'text'
                                    freely, set keep=0 to drop a row) and writes real_conversations.jsonl in the same schema as data/conversations.jsonl
                                    (user turns only), plus an aggregate summary.

Default private directory: ~/Documents/turing_real_data_private (mode 700). Override with --private-dir.
"""
import argparse, collections, csv, datetime as dt, glob, hashlib, json, os, random, re, sys, uuid

SEED = 20261024
NS = uuid.UUID("5a1f0f7e-3c0e-4d57-9e9a-2b7c1d0f4a11")
EXPORT_DIR = os.path.expanduser("~/Documents")
CHATGPT_GLOB = "chatgpt_personal_14.04.2026/conversations-*.json"
CLAUDE_FILE = "claude_max20_14.04.2026/conversations.json"
MIN_CHARS, MAX_CHARS = 20, 600
MAX_PER_CONVERSATION, MIN_PERSONAL_PER_CONVERSATION = 10, 3
QUOTAS = {"tr": 150, "en": 100, "mixed": 50}      # target messages per conversation language; the owner prunes in review

# ---------------------------------------------------------------- selection heuristics
TR_PERSONAL = re.compile(
    r"\b(ben|benim|bana|beni|bende|bence|annem|babam|kardeşim|kardeş|arkadaşım|eşim|sevgilim|abim|ablam|kızım|oğlum|dayım|teyzem|amcam|"
    r"patronum|ofisim|evim|işim|okulum|köpeğim|kedim)\b|\w+(ıyorum|iyorum|uyorum|üyorum|yorum|yorduk|dım|dim|dum|düm|tım|tim|tum|tüm|ım|im)\b", re.I)
EN_PERSONAL = re.compile(
    r"\b(I'm|I am|I was|I have|I had|I live|I work|I moved|I like|I love|I hate|I don't|I do not|I just|I will|I want|I used to|I'll|I've)\b|"
    r"\bmy (mom|dad|mother|father|wife|husband|girlfriend|boyfriend|friend|sister|brother|job|boss|dog|cat|apartment|house|flat|company|team|partner)\b")
TR_CHARS = re.compile(r"[çğıöşüÇĞİÖŞÜ]")
TR_WORDS = re.compile(r"\b(ve|bir|bu|ne|için|mi|mı|ama|çok|var|yok|ben|sen|nasıl|değil|gibi|daha|bana|şey|yani|tamam|evet|hayır)\b", re.I)
EN_WORDS = re.compile(r"\b(the|and|is|are|to|of|in|that|it|for|with|you|this|what|how|can|my|your|have|not)\b", re.I)
CODE_HINT = re.compile(r"```|Traceback|Exception|\bimport \w+|\bdef \w+\(|function\s*\(|=>|</?\w+>|\{\s*\"|^\s*[$#] ", re.M)
URL = re.compile(r"https?://\S+|www\.\S+", re.I)

# ---------------------------------------------------------------- masking
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
PHONE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{8,16}\d)(?!\d)")
LONG_DIGITS = re.compile(r"(?<!\d)\d{11,19}(?!\d)")
IBAN = re.compile(r"\b[A-Z]{2}\d{2}(?:[ ]?[A-Z0-9]{4}){3,7}(?:[ ]?[A-Z0-9]{1,4})?\b")
TOKEN = re.compile(r"\b(?:sk|pk|rk|ghp|gho|xox[a-z]|AKIA|AIza)[-_A-Za-z0-9]{12,}\b|\b[A-Za-z0-9+/_=-]{32,}\b")
# Only mask a value that follows a secret word with ":" or "=" (or "is"/"şu"-style cases are left to the owner's review): "kaç token" must stay readable.
SECRET_WORD = re.compile(r"((?:şifre|şifrem|sifre|parola|password|passwd|cvv|api key|apikey|secret)\w*\s*[:=]\s*)(\S+)", re.I)
ADDRESS = re.compile(r"\b[\w' ]{2,30}\s(?:Mah\.?|Mahallesi|Sok\.?|Sokak|Cad\.?|Cadde|Bulvar[ıi]?|Apt\.?|Apartman[ıi]?|Blok)\s*(?:No\s*[:.]?\s*\d+[\w/]*)?(?:\s*(?:D|Daire)\s*[:.]?\s*\d+)?", re.I)

TR_NAMES = ("Ahmet Mehmet Mustafa Ali Ayşe Fatma Emine Hatice Zeynep Elif Merve Büşra Esra Gamze Selin Deniz Ece Ezgi Burak Emre Cem Mert Kerem "
            "Onur Oğuz Umut Barış Berk Yusuf İbrahim Hasan Hüseyin Osman Murat Serkan Tolga Volkan Furkan Burcu Pınar Sibel Özge Derya Gizem Tuğçe "
            "İrem Melis Nur Sena Yağmur Ceren Aslı Damla Dilara Hande Kübra Beyza Betül Sevgi Seda Songül Nihal Leyla Cansu Defne Eda Gökhan "
            "Hakan Kaan Koray Levent Orhan Ozan Selçuk Taner Uğur Yasin Ömer Enes Eren Arda Doğan Batuhan Bilal Ceyda Dilek Feyza Gül Hülya İpek "
            "Kader Melek Nazlı Neslihan Nilay Özlem Rabia Sevim Tuba Yasemin Zehra Zübeyde Ahsen Metin Kemal Sinan Tuncay Ertan Recep Cengiz").split()
EN_NAMES = ("John James Michael David Robert William Richard Thomas Charles Sarah Emily Anna Maria Laura Chris Alex Tom Daniel Mark Paul Peter "
            "Jessica Jennifer Lisa Kate Emma Olivia Sophia Ben Sam Jack Josh Matt Nick Andrew Steve Kevin Brian Amy Rachel Hannah Megan Lauren").split()
FAKE_TR = ("Selim Aylin Tarık Nehir Okan Gonca Sevil Taylan Berna Cihan Mine Rüzgar Ilgaz Arzu Kemal Naz Orkun Şule Tamer Yeliz").split()
FAKE_EN = ("Tyler Nora Colin Maeve Owen Piper Dylan Wren Felix Iris Hugo Tessa Miles Quinn Eliot Sasha").split()
NAME_SET = {n: ("tr" if n in TR_NAMES else "en") for n in TR_NAMES + EN_NAMES}
CAP_TOKEN = re.compile(r"(?<![.!?]\s)(?<!^)\b([A-ZÇĞİÖŞÜ][a-zçğıöşü]{2,})\b")
ALLOW_CAPS = set(("Ocak Şubat Mart Nisan Mayıs Haziran Temmuz Ağustos Eylül Ekim Kasım Aralık Pazartesi Salı Çarşamba Perşembe Cuma Cumartesi Pazar "
                  "January February March April May June July August September October November December Monday Tuesday Wednesday Thursday Friday "
                  "Saturday Sunday Türkiye Turkey Türkçe Turkish English Almanya Germany Fransa France İstanbul Istanbul Ankara İzmir Izmir "
                  "Berlin London Paris Amsterdam ChatGPT Claude Python Swift Apple Google Android iOS iPhone Mac Excel Word").split())


def lang_of(text):
    tr = len(TR_CHARS.findall(text)) + 2 * len(TR_WORDS.findall(text))
    en = 2 * len(EN_WORDS.findall(text))
    if tr >= 2 and en >= 2: return "mixed"
    return "tr" if tr > en else "en"


def personal_score(text):
    return min(3, len(TR_PERSONAL.findall(text)) + len(EN_PERSONAL.findall(text)))


def usable(text):
    t = (text or "").strip()
    if not (MIN_CHARS <= len(t) <= MAX_CHARS) or CODE_HINT.search(t): return False
    stripped = URL.sub("", t)
    if len(stripped.strip()) < MIN_CHARS: return False
    if sum(c in "{}[]();=<>/\\_|*#" for c in t) / len(t) > 0.08: return False
    return len(t.split("\n")) <= 6


def pseudonym(name, salt):
    pool = FAKE_TR if NAME_SET[name] == "tr" else FAKE_EN
    return pool[int(hashlib.sha256((salt + name).encode()).hexdigest(), 16) % len(pool)]


def mask(text, salt):
    t = URL.sub("[URL]", text)
    t = EMAIL.sub("[EMAIL]", t)
    t = IBAN.sub("[IBAN]", t)
    t = SECRET_WORD.sub(lambda m: m.group(1) + "[SECRET]", t)
    t = TOKEN.sub("[TOKEN]", t)
    t = LONG_DIGITS.sub("[NUMBER]", t)
    t = PHONE.sub("[PHONE]", t)
    t = ADDRESS.sub("[ADDRESS]", t)
    t = re.sub(r"\b([A-ZÇĞİÖŞÜ][a-zçğıöşü]+)((?:'|’)\w+)?\b", lambda m: pseudonym(m.group(1), salt) + (m.group(2) or "") if m.group(1) in NAME_SET else m.group(0), t)
    return t


def check_tokens(masked):
    toks = {m.group(1) for m in CAP_TOKEN.finditer(masked)}
    return sorted(t for t in toks if t not in ALLOW_CAPS and t not in FAKE_TR and t not in FAKE_EN)


# ---------------------------------------------------------------- readers (user-authored text only)
def chatgpt_conversations():
    for f in sorted(glob.glob(os.path.join(EXPORT_DIR, CHATGPT_GLOB))):
        for c in json.load(open(f, encoding="utf-8")):
            if c.get("is_do_not_remember"): continue
            mapping, node, chain = c.get("mapping") or {}, c.get("current_node"), []
            while node:
                chain.append(node); node = (mapping.get(node) or {}).get("parent")
            msgs = []
            for nid in reversed(chain):
                m = (mapping.get(nid) or {}).get("message")
                if not m or m["author"]["role"] != "user": continue
                if m["content"].get("content_type") not in ("text", "multimodal_text"): continue
                parts = m["content"].get("parts") or []
                text = "\n".join(p for p in parts if isinstance(p, str))
                if text.strip(): msgs.append((m.get("create_time") or c.get("create_time"), text))
            yield "chatgpt", str(c.get("id") or c.get("conversation_id")), msgs


def claude_conversations():
    for c in json.load(open(os.path.join(EXPORT_DIR, CLAUDE_FILE), encoding="utf-8")):
        msgs = []
        for m in c.get("chat_messages", []):
            if m.get("sender") != "human": continue
            text = m.get("text") or ""
            if text.strip():
                ts = dt.datetime.fromisoformat(m["created_at"].replace("Z", "+00:00")).timestamp()
                msgs.append((ts, text))
        yield "claude", str(c.get("uuid")), sorted(msgs)


# ---------------------------------------------------------------- build
def build(private_dir):
    rng = random.Random(SEED)
    eligible = collections.defaultdict(list)      # language -> candidate runs
    seen_texts, totals = set(), collections.Counter()
    for source, conv_id, msgs in list(chatgpt_conversations()) + list(claude_conversations()):
        run = []
        for ts, text in msgs:
            totals[source + "_user_messages"] += 1
            if not usable(text): continue
            key = re.sub(r"\W+", " ", text.lower()).strip()
            if key in seen_texts: continue
            seen_texts.add(key)
            run.append((ts, text.strip()))
        personal = [r for r in run if personal_score(r[1]) >= 1]
        if len(personal) < MIN_PERSONAL_PER_CONVERSATION: continue
        # whole early run, so memory state can carry forward like in the app; cap length
        run = run[:MAX_PER_CONVERSATION]
        langs = collections.Counter(lang_of(t) for _, t in run)
        eligible[langs.most_common(1)[0][0]].append((source, conv_id, run))
        totals["eligible_conversations_" + langs.most_common(1)[0][0]] += 1
    chosen = []
    for lang, quota in QUOTAS.items():
        pool = eligible[lang]; rng.shuffle(pool); n = 0
        for item in pool:
            if n >= quota: break
            chosen.append((lang, item)); n += len(item[2])
    os.makedirs(private_dir, mode=0o700, exist_ok=True)
    rows, conv_counter = [], 0
    for lang, (source, conv_id, run) in chosen:
        conv_counter += 1
        cid = "r%03d" % conv_counter
        salt = hashlib.sha256((str(SEED) + conv_id).encode()).hexdigest()[:8]
        for i, (ts, text) in enumerate(run):
            masked = mask(text, salt)
            when = dt.datetime.fromtimestamp(float(ts), dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
            rows.append(dict(row_id=f"{cid}-{i:02d}", conversation=cid, source=source, lang=lang, created_at=when, text=masked,
                             check=" ".join(check_tokens(masked)), keep="1"))
    path = os.path.join(private_dir, "review.csv")
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    os.chmod(path, 0o600)
    print("aggregate only (no message text printed):")
    for k, v in sorted(totals.items()): print(f"  {k}: {v}")
    by = collections.Counter((r["source"], r["lang"]) for r in rows)
    print(f"selected {len(rows)} messages in {conv_counter} conversations;", dict(by))
    print(f"rows with check tokens: {sum(1 for r in rows if r['check'])}")
    print("review file:", path)


# ---------------------------------------------------------------- finalize
def to_xlsx(private_dir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    rows = list(csv.DictReader(open(os.path.join(private_dir, "review.csv"), encoding="utf-8")))
    wb = Workbook(); ws = wb.active; ws.title = "review"
    cols = ["row_id", "conversation", "source", "lang", "created_at", "text", "check", "keep"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1B3A5C")
    for r in rows:
        ws.append([r[c] for c in cols])
        cell = ws.cell(ws.max_row, 6); cell.data_type = "s"          # never let Excel read a message as a formula
        for c in ws[ws.max_row]: c.alignment = Alignment(wrap_text=True, vertical="top")
    for letter, w in zip("ABCDEFGH", (10, 12, 9, 7, 21, 90, 22, 7)): ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=False, showErrorMessage=True, error="1 = tut, 0 = sil")
    ws.add_data_validation(dv); dv.add(f"H2:H{len(rows) + 1}")
    out = os.path.join(private_dir, "review.xlsx")
    wb.save(out); os.chmod(out, 0o600)
    print("wrote", out, f"({len(rows)} rows)")


def parse_row_spec(spec):
    """'r002-03' (one row), 'r002-03..r002-09' (inclusive range) or 'r014' (the whole conversation) -> predicate on a row_id."""
    spec = spec.strip()
    if ".." in spec:
        lo, hi = (x.strip() for x in spec.split("..", 1))
        return lambda rid: lo <= rid <= hi
    if "-" in spec:
        return lambda rid: rid == spec
    return lambda rid: rid.startswith(spec + "-")


def apply_private_rules(rows, private_dir):
    """Two optional files the owner writes in the private directory (they never leave it):
    exclude_rows.txt  one row spec per line (see parse_row_spec); '#' starts a comment. These rows are dropped.
    mask_words.txt    one word or phrase per line (case-insensitive substring); every match becomes [REDACTED], e.g. the owner's
                      handle, full name, employer, school, town, product names that the generic masker cannot know."""
    ex_path, mw_path = os.path.join(private_dir, "exclude_rows.txt"), os.path.join(private_dir, "mask_words.txt")
    drop = []
    if os.path.exists(ex_path):
        drop = [parse_row_spec(l.split("#")[0]) for l in open(ex_path, encoding="utf-8") if l.split("#")[0].strip()]
    words = []
    if os.path.exists(mw_path):
        words = [l.strip() for l in open(mw_path, encoding="utf-8") if l.strip() and not l.lstrip().startswith("#")]
    kept, dropped, masked_hits = [], 0, 0
    for r in rows:
        if any(f(r["row_id"]) for f in drop):
            dropped += 1; continue
        for w in words:
            new = re.sub(re.escape(w), "[REDACTED]", r["text"], flags=re.I)
            masked_hits += new != r["text"]; r["text"] = new
        kept.append(r)
    print(f"private rules: {dropped} rows excluded, {len(words)} mask words applied ({masked_hits} replacements)")
    return kept


def finalize(private_dir):
    xlsx = os.path.join(private_dir, "review.xlsx")
    if os.path.exists(xlsx):
        from openpyxl import load_workbook
        sheet = load_workbook(xlsx)["review"]
        head = [c.value for c in sheet[1]]
        all_rows = [{h: ("" if v is None else str(v).strip() if h == "keep" else str(v)) for h, v in zip(head, [c.value for c in row])}
                    for row in sheet.iter_rows(min_row=2)]
        for r in all_rows:
            if r["keep"].endswith(".0"): r["keep"] = r["keep"][:-2]
    else:
        all_rows = list(csv.DictReader(open(os.path.join(private_dir, "review.csv"), encoding="utf-8")))
    rows = [r for r in all_rows if r["keep"].strip() not in ("0", "")]
    rows = apply_private_rules(rows, private_dir)
    convs = collections.OrderedDict()
    for r in rows:
        convs.setdefault(r["conversation"], []).append(r)
    out, n_msgs = [], 0
    for cid, items in convs.items():
        turns = []
        for r in items:
            turns.append({"id": str(uuid.uuid5(NS, r["row_id"])), "role": "user", "text": r["text"], "createdAt": r["created_at"]})
        n_msgs += len(turns)
        out.append({"conversation_id": cid, "persona": "real user (owner), masked", "lang": items[0]["lang"], "phenomena": [], "source": "real", "turns": turns})
    target = os.path.join(private_dir, "real_conversations.jsonl")
    with open(target, "w", encoding="utf-8") as f:
        for o in out: f.write(json.dumps(o, ensure_ascii=False) + "\n")
    os.chmod(target, 0o600)
    print(f"kept {n_msgs} messages in {len(out)} conversations -> {target}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("step", choices=["build", "xlsx", "finalize"])
    ap.add_argument("--private-dir", default=os.path.expanduser("~/Documents/turing_real_data_private"))
    a = ap.parse_args()
    {"build": build, "xlsx": to_xlsx, "finalize": finalize}[a.step](a.private_dir)
